"""
Word rapor oluşturma modülü
python-docx kullanarak Word belgeleri oluşturur
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from pathlib import Path
from typing import Dict, List, Optional
from datetime import datetime


class WordGenerator:
    """Word rapor oluşturucu sınıfı"""

    def __init__(self, config: Dict = None):
        """
        Args:
            config: Word konfigürasyonu
        """
        self.config = config or {}
        self.translations = self.config.get('translations', {})
        self.doc = Document()
        self._setup_styles()

        # Otokar renkleri
        self.colors = {
            'primary': RGBColor(0, 51, 102),  # #003366
            'secondary': RGBColor(255, 102, 0),  # #FF6600
            'success': RGBColor(0, 170, 68),  # #00AA44
            'danger': RGBColor(221, 0, 0),  # #DD0000
        }

    def _setup_styles(self):
        """Belge stillerini ayarla"""
        styles = self.doc.styles

        # Başlık stili
        if 'CustomTitle' not in styles:
            title_style = styles.add_style('CustomTitle', WD_STYLE_TYPE.PARAGRAPH)
            title_font = title_style.font
            title_font.name = 'Calibri'
            title_font.size = Pt(18)
            title_font.bold = True
            title_font.color.rgb = self.colors['primary']
            title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title_style.paragraph_format.space_after = Pt(12)

        # Heading stili
        if 'CustomHeading' not in styles:
            heading_style = styles.add_style('CustomHeading', WD_STYLE_TYPE.PARAGRAPH)
            heading_font = heading_style.font
            heading_font.name = 'Calibri'
            heading_font.size = Pt(14)
            heading_font.bold = True
            heading_font.color.rgb = self.colors['primary']
            heading_style.paragraph_format.space_before = Pt(12)
            heading_style.paragraph_format.space_after = Pt(6)

        # Subheading stili
        if 'CustomSubheading' not in styles:
            subheading_style = styles.add_style('CustomSubheading', WD_STYLE_TYPE.PARAGRAPH)
            subheading_font = subheading_style.font
            subheading_font.name = 'Calibri'
            subheading_font.size = Pt(12)
            subheading_font.bold = True
            subheading_font.color.rgb = self.colors['secondary']
            subheading_style.paragraph_format.space_before = Pt(8)
            subheading_style.paragraph_format.space_after = Pt(4)

    def create_cover_page(
        self,
        title: str,
        test_info: Dict,
        logo_path: Optional[Path] = None
    ):
        """
        Kapak sayfası oluştur

        Args:
            title: Rapor başlığı
            test_info: Test bilgileri
            logo_path: Logo dosya yolu
        """
        # Logo
        if logo_path and Path(logo_path).exists():
            try:
                para = self.doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run()
                run.add_picture(str(logo_path), width=Inches(3))
            except:
                pass

        # Boşluk
        self.doc.add_paragraph()

        # Başlık
        title_para = self.doc.add_paragraph(title)
        title_para.style = 'CustomTitle'

        # Boşluk
        self.doc.add_paragraph()

        # Test bilgileri
        test_type = test_info.get('test_type', 'N/A')
        test_date = test_info.get('test_date', datetime.now().strftime('%Y-%m-%d'))
        vehicle = test_info.get('vehicle_type', 'N/A')

        # Tablo
        table = self.doc.add_table(rows=3, cols=2)
        table.style = 'Light Grid Accent 1'

        table.cell(0, 0).text = self.translate('test_type')
        table.cell(0, 1).text = test_type
        table.cell(1, 0).text = self.translate('test_date')
        table.cell(1, 1).text = test_date
        table.cell(2, 0).text = self.translate('vehicle')
        table.cell(2, 1).text = vehicle

        # Başlık sütununu kalın yap
        for i in range(3):
            table.cell(i, 0).paragraphs[0].runs[0].font.bold = True

        # Boşluk
        self.doc.add_paragraph()
        self.doc.add_paragraph()

        # Rapor tarihi
        report_date = datetime.now().strftime('%Y-%m-%d %H:%M')
        date_para = self.doc.add_paragraph(
            f"{self.translate('generated_on')}: {report_date}"
        )
        date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Sayfa sonu
        self.doc.add_page_break()

    def add_section_title(self, title: str):
        """Bölüm başlığı ekle"""
        heading = self.doc.add_paragraph(title)
        heading.style = 'CustomHeading'

    def add_subsection_title(self, title: str):
        """Alt bölüm başlığı ekle"""
        subheading = self.doc.add_paragraph(title)
        subheading.style = 'CustomSubheading'

    def add_paragraph(self, text: str, bold: bool = False):
        """
        Paragraf ekle

        Args:
            text: Metin
            bold: Kalın yazılsın mı
        """
        para = self.doc.add_paragraph()
        run = para.add_run(text)
        if bold:
            run.font.bold = True

    def add_statistics_table(
        self,
        stats_data: Dict[str, Dict],
        title: Optional[str] = None
    ):
        """
        İstatistik tablosu ekle

        Args:
            stats_data: İstatistik verileri {metrik: {stat: value}}
            title: Tablo başlığı
        """
        if title:
            self.add_subsection_title(title)

        # Başlık satırı
        headers = [self.translate('parameter')]
        if stats_data:
            first_key = list(stats_data.keys())[0]
            stat_keys = list(stats_data[first_key].keys())
            headers.extend([self.translate(k) if k in ['mean', 'min', 'max', 'std', 'median']
                           else k.capitalize() for k in stat_keys])

        # Tablo oluştur
        num_rows = len(stats_data) + 1
        num_cols = len(headers)
        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.style = 'Light Grid Accent 1'

        # Başlık satırı
        for i, header in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = header
            # Başlık formatı
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            # Arka plan rengi (şading)
            shading_elm = self._get_or_create_shading(cell)
            shading_elm.set(qn('w:fill'), '003366')

        # Veri satırları
        row_idx = 1
        for param, stats in stats_data.items():
            table.cell(row_idx, 0).text = param
            col_idx = 1
            for value in stats.values():
                if isinstance(value, (int, float)):
                    table.cell(row_idx, col_idx).text = f"{value:.2f}"
                else:
                    table.cell(row_idx, col_idx).text = str(value)
                col_idx += 1
            row_idx += 1

        self.doc.add_paragraph()

    def add_comparison_table(
        self,
        comparison_data: List[Dict],
        columns: List[str],
        title: Optional[str] = None
    ):
        """
        Karşılaştırma tablosu ekle

        Args:
            comparison_data: Karşılaştırma verileri
            columns: Sütun isimleri
            title: Tablo başlığı
        """
        if title:
            self.add_subsection_title(title)

        # Tablo oluştur
        num_rows = len(comparison_data) + 1
        num_cols = len(columns)
        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.style = 'Light Grid Accent 1'

        # Başlık satırı
        for i, col in enumerate(columns):
            cell = table.cell(0, i)
            cell.text = self.translate(col)
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            shading_elm = self._get_or_create_shading(cell)
            shading_elm.set(qn('w:fill'), 'FF6600')

        # Veri satırları
        for row_idx, row_dict in enumerate(comparison_data, start=1):
            for col_idx, col in enumerate(columns):
                value = row_dict.get(col, '-')
                table.cell(row_idx, col_idx).text = str(value)

        self.doc.add_paragraph()

    def add_requirement_check_table(
        self,
        check_results: List[Dict],
        title: Optional[str] = None
    ):
        """
        Gereksinim kontrolü tablosu ekle

        Args:
            check_results: Kontrol sonuçları
            title: Tablo başlığı
        """
        if title:
            self.add_subsection_title(title)

        # Tablo oluştur
        num_rows = len(check_results) + 1
        table = self.doc.add_table(rows=num_rows, cols=4)
        table.style = 'Light Grid Accent 1'

        # Başlık satırı
        headers = [
            self.translate('parameter'),
            self.translate('requirement'),
            self.translate('actual'),
            self.translate('status')
        ]
        for i, header in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = header
            cell.paragraphs[0].runs[0].font.bold = True
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            shading_elm = self._get_or_create_shading(cell)
            shading_elm.set(qn('w:fill'), '003366')

        # Veri satırları
        for row_idx, result in enumerate(check_results, start=1):
            param = result.get('parameter', '-')
            req = result.get('requirement', '-')
            actual = result.get('actual', '-')
            status = result.get('status', 'UNKNOWN')

            # Format
            if isinstance(req, (int, float)):
                req = f"{req:.2f}"
            if isinstance(actual, (int, float)):
                actual = f"{actual:.2f}"

            table.cell(row_idx, 0).text = param
            table.cell(row_idx, 1).text = str(req)
            table.cell(row_idx, 2).text = str(actual)

            # Durum hücresi - renklendirme
            status_cell = table.cell(row_idx, 3)
            status_cell.text = self.translate(status.lower())
            status_para = status_cell.paragraphs[0]
            status_run = status_para.runs[0]
            status_run.font.bold = True

            if status == 'PASS':
                status_run.font.color.rgb = RGBColor(255, 255, 255)
                shading_elm = self._get_or_create_shading(status_cell)
                shading_elm.set(qn('w:fill'), '00AA44')
            elif status == 'FAIL':
                status_run.font.color.rgb = RGBColor(255, 255, 255)
                shading_elm = self._get_or_create_shading(status_cell)
                shading_elm.set(qn('w:fill'), 'DD0000')

        self.doc.add_paragraph()

    def add_image(
        self,
        image_path: Path,
        width: float = 6,
        title: Optional[str] = None
    ):
        """
        Görsel ekle

        Args:
            image_path: Görsel dosya yolu
            width: Genişlik (inch)
            title: Görsel başlığı
        """
        if title:
            self.add_paragraph(title, bold=True)

        try:
            para = self.doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(str(image_path), width=Inches(width))
            self.doc.add_paragraph()
        except Exception as e:
            self.add_paragraph(f"[Görsel yüklenemedi: {e}]")

    def add_page_break(self):
        """Sayfa sonu ekle"""
        self.doc.add_page_break()

    def translate(self, key: str) -> str:
        """Metni çevir"""
        return self.translations.get(key, key.replace('_', ' ').title())

    def _get_or_create_shading(self, cell):
        """Hücre shading elementi al veya oluştur"""
        tcPr = cell._element.get_or_add_tcPr()
        shading = tcPr.find(qn('w:shd'))
        if shading is None:
            shading = self.doc._element.makeelement(qn('w:shd'))
            tcPr.append(shading)
        return shading

    def save(self, output_path: Path):
        """
        Word belgesini kaydet

        Args:
            output_path: Çıktı dosya yolu
        """
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        self.doc.save(str(output_path))
        print(f"✓ Word raporu oluşturuldu: {output_path}")

        return output_path
